"""OPS HUB Toilet QR request functions with optional Telegram delivery."""
import hashlib
import html
import io
import secrets
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import qrcode
import os
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from db import get_conn, ph, DB_TYPE
from master_loader import load_toilet_request_rows

SG = ZoneInfo('Asia/Singapore')
REQUESTS = [
    'Top up toilet paper', 'Top up hand towels', 'Top up toilet seat sanitiser',
    'Top up hand soap', 'Top up sanitary bags', 'Clear trash',
    'Clean wet floor', 'Clean stained toilet seat', 'Clean dirty floor',
    'Toilet odour', 'Pests', 'Other Issues',
]

def now_sg():
    return datetime.now(SG).replace(tzinfo=None)

def route_team(toilet, when=None):
    when = when or now_sg()
    return toilet['supervisor_group'] if str(toilet.get('zone') or '').strip().lower() == 'demo' else ('MN' if when.hour >= 21 or when.hour < 7 else toilet['supervisor_group'])

def get_toilet(token):
    digest = hashlib.sha256(token.encode()).hexdigest()
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT toilet_id FROM toilet_qr_tokens WHERE token_hash={ph()} AND active={ph()}', (digest, True if DB_TYPE == 'postgres' else 1))
        row = cur.fetchone()
    finally:
        conn.close()
    if not row: return None
    return next((t for t in load_toilet_request_rows() if t['toilet_id'] == row['toilet_id']), None)

def ensure_tokens():
    """Generate missing QR tokens, returning only newly generated secrets."""
    new = {}
    conn = get_conn()
    try:
        cur = conn.cursor()
        for t in load_toilet_request_rows():
            cur.execute(f'SELECT toilet_id FROM toilet_qr_tokens WHERE toilet_id={ph()}', (t['toilet_id'],))
            if cur.fetchone(): continue
            raw = secrets.token_urlsafe(24)
            cur.execute(f'INSERT INTO toilet_qr_tokens (toilet_id,token_hash) VALUES ({ph()},{ph()})', (t['toilet_id'], hashlib.sha256(raw.encode()).hexdigest()))
            new[t['toilet_id']] = raw
        conn.commit()
    except Exception:
        conn.rollback(); raise
    finally: conn.close()
    return new

def submit_many(toilet, request_types, details=''):
    """Save new issues atomically, then send one notification for the batch."""
    chosen = list(dict.fromkeys(request_types))
    allowed = request_options(toilet)
    if not chosen or any(item not in allowed for item in chosen):
        raise ValueError('Select at least one valid issue.')
    details = details.strip() if 'Other Issues' in chosen else ''
    if 'Other Issues' in chosen and not details:
        raise ValueError('Please describe the other issue.')
    if len(details) > 200:
        raise ValueError('Maximum 200 characters.')
    stamp = now_sg()
    team = route_team(toilet, stamp)
    cutoff = (stamp - timedelta(minutes=15)).strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    ids, accepted = [], []
    try:
        cur = conn.cursor()
        # Serialise submissions for this toilet on PostgreSQL so concurrent visitors
        # cannot both pass the duplicate check for the same issue.
        if DB_TYPE == 'postgres':
            cur.execute(f'SELECT toilet_id FROM toilet_qr_tokens WHERE toilet_id={ph()} FOR UPDATE',
                        (toilet['toilet_id'],))
        for item in chosen:
            item_details = details if item == 'Other Issues' else ''
            cur.execute(f"""SELECT request_id FROM toilet_requests WHERE toilet_id={ph()}
                AND request_type={ph()} AND COALESCE(details,'')={ph()}
                AND created_at >= {ph()} LIMIT 1""",
                (toilet['toilet_id'], item, item_details, cutoff))
            if cur.fetchone():
                continue
            params = (toilet['toilet_id'], toilet['toilet_name'], toilet['tower'],
                      toilet['level'], item, item_details, team,
                      stamp.strftime('%Y-%m-%d %H:%M:%S'))
            marks = ','.join([ph()] * 8)
            suffix = ' RETURNING request_id' if DB_TYPE == 'postgres' else ''
            cur.execute(f"""INSERT INTO toilet_requests
                (toilet_id,toilet_name,tower,level,request_type,details,routed_team,created_at)
                VALUES ({marks}){suffix}""", params)
            ids.append(cur.fetchone()['request_id'] if DB_TYPE == 'postgres' else cur.lastrowid)
            accepted.append((item, item_details))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    if not ids:
        return {'ids': [], 'skipped': len(chosen), 'status': 'DUPLICATE'}

    # All requests are already committed. Notification failures cannot undo them.
    try:
        from toilet_telegram import send_many as telegram_send_many
        result = telegram_send_many(toilet, accepted, stamp)
        status = result.get('status', 'FAILED')
    except Exception:
        status = 'FAILED'
    try:
        conn = get_conn()
        try:
            cur = conn.cursor()
            for request_id in ids:
                cur.execute(f'UPDATE toilet_requests SET notification_status={ph()} WHERE request_id={ph()}',
                            (status, request_id))
            conn.commit()
        except Exception:
            conn.rollback()
        finally:
            conn.close()
    except Exception:
        pass
    return {'ids': ids, 'skipped': len(chosen) - len(ids), 'status': status}


def submit(toilet, request_type, details=''):
    """Backward-compatible single-issue entry point."""
    result = submit_many(toilet, [request_type], details)
    return result['ids'][0] if result['ids'] else None

def recent():
    conn = get_conn()
    try:
        cur = conn.cursor(); cur.execute('SELECT * FROM toilet_requests ORDER BY request_id DESC LIMIT 20')
        return [dict(r) for r in cur.fetchall()]
    finally: conn.close()

def export_report():
    conn = get_conn()
    try:
        cur = conn.cursor();cur.execute('SELECT * FROM toilet_requests ORDER BY created_at,request_id')
        rows = [dict(r) for r in cur.fetchall()]
    finally: conn.close()
    wb = Workbook(); wb.remove(wb.active)
    cols = ['request_id','created_at','toilet_id','toilet_name','tower','level','request_type','details','routed_team','notification_status']
    demo_ids = {t['toilet_id'] for t in load_toilet_request_rows()
                if str(t.get('zone') or '').strip().lower() == 'demo'}
    for r in rows:
        if r['toilet_id'] in demo_ids:
            continue
        month = str(r['created_at'])[:7]
        name = datetime.strptime(month,'%Y-%m').strftime('%b %Y')
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        if ws.max_row == 1 and ws['A1'].value is None:
            ws.append(cols)
            for cell in ws[1]: cell.font = Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='17365D')
            ws.freeze_panes='A2';ws.auto_filter.ref='A1:J1'
        ws.append([str(r[c]) if r[c] is not None else '' for c in cols])
    if not wb.sheetnames: wb.create_sheet('No Requests').append(['No operational requests recorded'])
    for ws in wb:
        for column in ws.columns:
            letter=column[0].column_letter
            ws.column_dimensions[letter].width=min(48,max(14,max(len(str(c.value or '')) for c in column)+2))
    output=io.BytesIO();wb.save(output);return output.getvalue()

# Presentation labels and separately maintained illustration assets.
REQUEST_DISPLAY = {
    'Top up toilet paper': ('paper', 'Toilet paper'),
    'Top up hand towels': ('towels', 'Hand towels'),
    'Top up toilet seat sanitiser': ('sanitiser', 'Seat sanitiser'),
    'Top up hand soap': ('soap', 'Hand soap'),
    'Top up sanitary bags': ('bags', 'Sanitary bags'),
    'Clear trash': ('bin', 'Clear trash bin'),
    'Clean wet floor': ('wet', 'Wet floor'),
    'Clean stained toilet seat': ('seat', 'Stained toilet seat'),
    'Clean dirty floor': ('floor', 'Dirty floor'),
    'Toilet odour': ('odour', 'Toilet odour'),
    'Pests': ('pests', 'Pests spotted'),
    'Other Issues': ('other', 'Other issues'),
}

def request_options(toilet):
    """Sanitary bags: female/unisex only; never male/accessible."""
    toilet_type = str(toilet.get('toilet_type') or '').strip().lower()
    return [item for item in REQUESTS
            if item != 'Top up sanitary bags' or toilet_type in ('female', 'unisex')]


def _public_style():
    """Render consistent light mobile styling; icon art is loaded from local SVG assets."""
    from pathlib import Path
    import base64
    asset_dir = Path(__file__).resolve().parent / 'housekeeping_request_assets' / 'icons'
    icons_css = []
    for idx, key in enumerate(REQUEST_DISPLAY):
        asset = asset_dir / (REQUEST_DISPLAY[key][0] + '.svg')
        if asset.is_file():
            encoded = base64.b64encode(asset.read_bytes()).decode('ascii')
            selector = '.st-key-hk_other button' if key == 'Other Issues' else f'.st-key-hk_option_{idx} button'
            icons_css.append(
                f'{selector}::before {{content:"";display:block;'
                f'width:44px;height:44px;margin:0 auto 6px;'
                f'background:center / contain no-repeat url("data:image/svg+xml;base64,{encoded}");}}'
            )
    css = """
    <style>
    /* Public-only page: deliberately override Streamlit dark theme. */
    [data-testid="stAppViewContainer"], [data-testid="stMain"],
    [data-testid="stMainBlockContainer"] {background:#fff !important;color:#10375b !important;}
    /* Hide Streamlit's public-facing toolbar and menu. */
    [data-testid="stHeader"], [data-testid="stToolbar"],
    [data-testid="stDecoration"], [data-testid="stStatusWidget"],
    #MainMenu, footer {display:none !important;}
    [data-testid="stAppViewContainer"] {overflow-x:hidden !important;}
    [data-testid="stMainBlockContainer"] {
        padding:1.35rem 0.9rem 2rem !important;
        max-width:560px !important;
        min-width:0 !important;
        box-sizing:border-box !important;
    }
    [data-testid="stMain"], [data-testid="stMainBlockContainer"],
    [data-testid="stMainBlockContainer"] > div {
        min-width:0 !important;max-width:100% !important;
    }
    [data-testid="stMain"] h1,[data-testid="stMain"] h2,
    [data-testid="stMain"] h3,[data-testid="stMain"] p {color:#10375b;}
    .st-key-hk_header {text-align:center;padding:16px 4px 10px;overflow:visible !important;
        background:linear-gradient(130deg,#fff 65%,#e8f8e7);border-radius:16px;}
    .st-key-hk_header h1 {color:#10375b !important;font-size:clamp(1.5rem,5.2vw,2.1rem);
        line-height:1.35 !important;margin:0 !important;padding:3px 0 2px !important;font-weight:800;}
    .st-key-hk_header [data-testid="stMarkdownContainer"] {overflow:visible !important;}
    .st-key-hk_header p {color:#31516a !important;margin:7px 0 2px;font-size:.95rem;}
    .st-key-hk_location {background:#f1f5f8;border:1px solid #dfe7ed;
        border-radius:14px;padding:15px 16px;margin:2px 0 5px;color:#10375b;
        min-height:88px;display:flex;flex-direction:column;justify-content:center;
        box-sizing:border-box;}
    .st-key-hk_location [data-testid="stMarkdownContainer"] {width:100%;}
    .st-key-hk_location p {margin:0 !important;color:#10375b !important;
        line-height:1.4 !important;overflow-wrap:anywhere;}
    .st-key-hk_location .hk-location-title {font-weight:750;margin-bottom:6px;}
    .st-key-hk_location .hk-location-name {font-size:.94rem;}
    /* Streamlit stacks columns on narrow screens by default. Force a real,
       shrinkable two-column grid without the oversized min-width columns. */
    .st-key-hk_options, .st-key-hk_options > div,
    .st-key-hk_options [data-testid="stVerticalBlock"],
    .st-key-hk_options [data-testid="stHorizontalBlock"] {
        width:100% !important;max-width:100% !important;min-width:0 !important;
        box-sizing:border-box !important;
    }
    .st-key-hk_options [data-testid="stHorizontalBlock"] {
        display:grid !important;
        grid-template-columns:repeat(2,minmax(0,1fr)) !important;
        gap:10px !important;align-items:stretch !important;
    }
    .st-key-hk_options [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] {
        width:100% !important;min-width:0 !important;max-width:100% !important;
        flex:none !important;box-sizing:border-box !important;
    }
    .st-key-hk_options [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] > div {
        width:100% !important;min-width:0 !important;
    }
    .st-key-hk_options [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] [data-testid="stVerticalBlock"] {
        gap:0 !important;
    }
    .st-key-hk_options div[data-testid="stButton"] button {
        display:block !important;min-height:104px !important;width:100% !important;
        border:1.5px solid #dce6ed !important;border-radius:14px !important;
        background:#fff !important;color:#10375b !important;
        font-weight:650 !important;white-space:normal !important;
        line-height:1.25 !important;font-size:.91rem !important;
        padding:9px 5px !important;box-shadow:0 2px 6px #0d30520c !important;
    }
    .st-key-hk_options div[data-testid="stButton"] button p {
        color:inherit !important;font-size:inherit !important;
        margin:0 !important;overflow-wrap:anywhere !important;}
    .st-key-hk_options div[data-testid="stButton"] button[kind="primary"] {
        background:#e4f7e9 !important;border:2px solid #2f9f59 !important;
        color:#10375b !important;}
    .st-key-hk_other div[data-testid="stButton"] button {
        min-height:72px !important;width:100% !important;
        border:1.5px solid #dce6ed !important;border-radius:14px !important;
        background:#fff !important;color:#10375b !important;
        font-weight:650 !important;white-space:normal !important;
    }
    .st-key-hk_other div[data-testid="stButton"] button[kind="primary"] {
        background:#e4f7e9 !important;border:2px solid #2f9f59 !important;
    }
    .st-key-hk_other div[data-testid="stButton"] button p {color:inherit !important;}
    .st-key-hk_submit div[data-testid="stButton"] button {
        min-height:54px !important;border-radius:12px !important;
        background:#103b60 !important;color:#fff !important;
        font-weight:700 !important;border:0 !important;width:100% !important;}
    .st-key-hk_submit div[data-testid="stButton"] button *,
    .st-key-hk_submit div[data-testid="stButton"] button p,
    .st-key-hk_submit div[data-testid="stButton"] button span {
        color:#fff !important;-webkit-text-fill-color:#fff !important;opacity:1 !important;}
    .st-key-hk_submit div[data-testid="stButton"] button:disabled {
        opacity:1 !important;background:#8aa5ba !important;color:#fff !important;}
    .st-key-hk_submit div[data-testid="stButton"] button:disabled * {
        color:#fff !important;-webkit-text-fill-color:#fff !important;opacity:1 !important;}
    .st-key-hk_success {text-align:center;padding:14px 8px 4px;}
    .st-key-hk_success .hk-success-hero {padding:25px 0 15px;}
    .st-key-hk_success .hk-check-circle {position:relative;width:124px;height:124px;
        border-radius:50%;background:#c6efb9;margin:0 auto 20px;display:flex;
        align-items:center;justify-content:center;}
    .st-key-hk_success .hk-check-circle svg {width:76px;height:76px;}
    .st-key-hk_success .hk-spark {position:absolute;color:#42b35c;font-size:27px;}
    .st-key-hk_success .hk-spark-left {left:-28px;bottom:4px;}
    .st-key-hk_success .hk-spark-right {right:-28px;top:4px;}
    .st-key-hk_success h2 {color:#10375b !important;font-size:clamp(1.65rem,6vw,2rem);
        margin:0 0 12px !important;text-align:center;}
    .st-key-hk_success p {color:#31516a !important;text-align:center;
        line-height:1.5;margin:0 auto 24px !important;max-width:330px;}
    .st-key-hk_success .hk-wait {display:flex;align-items:center;gap:14px;
        background:#f3f5f7;border-radius:14px;padding:17px 16px;
        text-align:left;color:#10375b;max-width:440px;margin:0 auto 22px;}
    .st-key-hk_success .hk-wait svg {width:32px;height:32px;flex-shrink:0;}
    .st-key-hk_success .hk-wait span {font-size:.91rem;line-height:1.5;}
    .st-key-hk_back div[data-testid="stButton"] button {background:#103b60 !important;
        color:#fff !important;border:0 !important;border-radius:12px !important;
        min-height:52px !important;font-weight:700 !important;width:100% !important;}
    .st-key-hk_back div[data-testid="stButton"] button * {
        color:#fff !important;-webkit-text-fill-color:#fff !important;}
    @media(max-width:380px) {
        .st-key-hk_options div[data-testid="stButton"] button {
            min-height:108px !important;font-size:.83rem !important;padding:9px 4px !important;}
        .st-key-hk_header h1 {font-size:1.48rem !important;}
    }
    </style>
    """
    st.markdown(css.replace('</style>', '\n' + '\n'.join(icons_css) + '\n</style>'),
                unsafe_allow_html=True)


def public_page(token):
    toilet = get_toilet(token)
    if not toilet:
        st.error('Invalid or inactive toilet QR code.')
        return
    _public_style()
    with st.container(key='hk_header'):
        st.markdown('<h1>✦ Housekeeping Request</h1>'
                    '<p>Help us keep this restroom<br>clean and comfortable</p>',
                    unsafe_allow_html=True)
    with st.container(key='hk_location'):
        tower = html.escape(str(toilet['tower']))
        level = html.escape(str(toilet['level']))
        name = html.escape(str(toilet['toilet_name']))
        st.markdown(
            f'<div class="hk-location-title">📍 Tower {tower} · Level {level}</div>'
            f'<div class="hk-location-name">{name}</div>',
            unsafe_allow_html=True,
        )

    key = 'toilet_last_' + token
    last = st.session_state.get(key)
    if last and (now_sg() - last).total_seconds() < 900:
        if not st.session_state.get('hk_back_' + token, False):
            with st.container(key='hk_success'):
                st.markdown('''
                <div class="hk-success-hero">
                  <div class="hk-check-circle">
                    <span class="hk-spark hk-spark-left">✦</span>
                    <svg viewBox="0 0 80 80" aria-hidden="true">
                      <path d="M15 40 L32 57 L66 22" fill="none" stroke="#10375b"
                        stroke-width="10" stroke-linecap="round" stroke-linejoin="round"/>
                    </svg>
                    <span class="hk-spark hk-spark-right">✦</span>
                  </div>
                </div>
                <h2>Request received!</h2>
                <p>Thank you for helping us keep this restroom clean and comfortable.</p>
                <div class="hk-wait">
                  <svg viewBox="0 0 32 32" fill="none" stroke="#10375b"
                       stroke-width="2.4" stroke-linecap="round" aria-hidden="true">
                    <circle cx="16" cy="16" r="12"/><path d="M16 8v9l6 4"/>
                  </svg>
                  <span>Please wait <strong>15 minutes</strong> before submitting another request from this page.</span>
                </div>
                ''', unsafe_allow_html=True)
            if st.session_state.get('hk_skipped_' + token, 0):
                st.caption(f"{st.session_state['hk_skipped_' + token]} previously reported issue(s) were not resubmitted.")
            with st.container(key='hk_back'):
                if st.button('Back to start', key='hk_back_button_' + token,
                             use_container_width=True):
                    st.session_state['hk_back_' + token] = True
                    st.rerun()
        else:
            st.info('Thank you! Please wait 15 minutes before submitting another request from this page.')
        return
    st.session_state.pop('hk_back_' + token, None)
    st.session_state.pop('hk_skipped_' + token, None)

    st.markdown('### What needs attention?')
    st.caption('Select all issues that need attention.')
    options = request_options(toilet)
    selected_key = 'toilet_selected_' + token
    selected = st.session_state.get(selected_key, [])
    if not isinstance(selected, list):
        selected = [selected] if selected else []
    selected = [item for item in selected if item in options]
    st.session_state[selected_key] = selected
    grid_options = [o for o in options if o != 'Other Issues']
    with st.container(key='hk_options'):
        for idx in range(0, len(grid_options), 2):
            cols = st.columns(2, gap='small')
            for offset, col in enumerate(cols):
                if idx + offset >= len(grid_options):
                    continue
                option = grid_options[idx + offset]
                _icon_name, label = REQUEST_DISPLAY[option]
                with col:
                    with st.container(key=f'hk_option_{REQUESTS.index(option)}'):
                        clicked = st.button(label, key=f'hk_{token}_{idx+offset}',
                                 type='primary' if option in selected else 'secondary',
                                 use_container_width=True)
                    if clicked:
                        st.session_state[selected_key] = (
                            [o for o in selected if o != option] if option in selected
                            else selected + [option])
                        st.rerun()

    if 'Other Issues' in options:
        with st.container(key='hk_other'):
            clicked = st.button('Other issues', key=f'hk_{token}_other',
                                type='primary' if 'Other Issues' in selected else 'secondary',
                                use_container_width=True)
            if clicked:
                st.session_state[selected_key] = (
                    [o for o in selected if o != 'Other Issues'] if 'Other Issues' in selected
                    else selected + ['Other Issues'])
                st.rerun()

    selected = st.session_state.get(selected_key, [])
    details = ''
    if 'Other Issues' in selected:
        details = st.text_area('Describe the issue', max_chars=200,
                               key='hk_details_' + token,
                               placeholder='Tell us what needs attention…')
    count = len(selected)
    label = f'Submit {count} issue' + ('s' if count != 1 else '') if count else 'Submit Request'
    with st.container(key='hk_submit'):
        pressed = st.button(label, disabled=not bool(selected),
                            use_container_width=True, key='hk_submit_' + token)
    if pressed:
        try:
            result = submit_many(toilet, selected, details)
            if result['ids']:
                st.session_state[key] = now_sg()
                st.session_state[selected_key] = []
                st.session_state['hk_skipped_' + token] = result['skipped']
                st.rerun()
            else:
                st.info('All selected issues have already been reported recently. Thank you.')
        except ValueError as exc:
            st.warning(str(exc))
        except Exception:
            st.error('Unable to submit. Please try again later.')

def admin_page():
    if st.session_state.get('user',{}).get('role')!='ADMIN':
        st.error('Administrator access required.');return
    st.title('Toilet Requests')
    st.caption('Telegram notification is controlled by TOILET_TELEGRAM_MODE. Requests are saved even if delivery fails.')
    st.subheader('Latest 20 requests')
    rows=recent()
    if rows: st.dataframe(rows,hide_index=True,use_container_width=True)
    else: st.info('No requests yet.')
    st.download_button('Generate Report',export_report(),file_name='Toilet_Requests.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    st.divider()
    st.subheader('Generate QR codes')
    st.warning('QR tokens are shown only when first generated. Save the ZIP immediately; do not share it publicly.')
    default_url = os.environ.get("PUBLIC_APP_URL", "")
    base = st.text_input(
        'Base URL for new QR codes',
        value=default_url,
        placeholder='http://192.168.1.100:8501',
        help='For local testing, enter this computer’s LAN address (not localhost). Your phone must be on the same network.',
    ).strip()
    if st.button('Generate missing toilet QR codes'):
        from urllib.parse import urlsplit
        parsed = urlsplit(base)
        if parsed.scheme not in ('http', 'https') or not parsed.netloc or parsed.username or parsed.password or parsed.query or parsed.fragment:
            st.error('Enter a valid base URL, e.g. http://192.168.1.100:8501 (without query parameters).')
        else:
            new = ensure_tokens()
            if not new:
                st.info('All toilets already have tokens. Existing QR codes remain valid.')
            else:
                import zipfile
                z = io.BytesIO()
                with zipfile.ZipFile(z, 'w') as archive:
                    for toilet_id, raw in new.items():
                        url = f"{base.rstrip('/')}?toilet_qr={raw}"
                        qr = qrcode.make(url)
                        img = io.BytesIO()
                        qr.save(img, format='PNG')
                        archive.writestr(f'{toilet_id}.png', img.getvalue())
                        archive.writestr(f'{toilet_id}.txt', url)
                st.session_state['toilet_new_qr_zip'] = z.getvalue()
                st.success(f'Generated {len(new)} new toilet QR code(s). Download the ZIP now.')
    if st.session_state.get('toilet_new_qr_zip'):
        st.download_button(
            'Download new QR ZIP',
            data=st.session_state['toilet_new_qr_zip'],
            file_name='Toilet_QR_New.zip',
            mime='application/zip',
        )
    st.divider()
    st.subheader('A4 Housekeeping Request Posters')
    st.caption('Upload the ORIGINAL toilet QR ZIP (PNG + TXT pairs). Existing token hashes cannot recreate the original links. Posters do not change any QR tokens.')
    poster_zip = st.file_uploader('Original Toilet QR ZIP', type=['zip'], key='toilet_poster_zip')
    if poster_zip is not None:
        try:
            from toilet_poster import parse_qr_zip, build_pdf
            qr_items = parse_qr_zip(poster_zip.getvalue())
            toilets = load_toilet_request_rows()
            eligible = [t for t in toilets if str(t.get('toilet_id')) in qr_items and str(t.get('active', 'Y')).upper() not in ('N', 'NO', 'FALSE', '0')]
            if not eligible:
                st.warning('No matching active toilets found. Use the original ZIP generated for these toilets.')
            else:
                st.success(f'{len(eligible)} active toilet(s) matched their original QR images.')
                labels = {str(t['toilet_id']): f"{t['toilet_id']} — {t.get('toilet_name') or t.get('name') or ''}" for t in eligible}
                chosen = st.selectbox('Single poster', list(labels), format_func=lambda k: labels[k], key='toilet_poster_single')
                single_pdf, _ = build_pdf(qr_items, toilets, selected={chosen})
                st.download_button('Download Single A4 Poster', single_pdf, file_name=f'Housekeeping_Poster_{chosen}.pdf', mime='application/pdf')
                all_pdf, ids = build_pdf(qr_items, toilets)
                st.download_button(f'Download All A4 Posters ({len(ids)})', all_pdf, file_name='Housekeeping_Posters_All.pdf', mime='application/pdf')
        except Exception as e:
            st.error(f'Poster generation failed: {e}')
    if st.button('Back to Home'):
        st.session_state['page']='home';st.rerun()
